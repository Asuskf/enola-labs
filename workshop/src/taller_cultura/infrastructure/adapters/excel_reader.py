"""Adaptador de salida: lee el Excel original del taller de cultura.

Implementa el puerto `LectorTaller`. Es el único lugar del proyecto que
conoce la disposición física, irregular, del archivo `.xlsx` original
(hojas `CALIFICADORES` y `TALLER`). Si el formato de origen cambia, solo
este archivo debería tocarse.

Estructura real de la hoja TALLER (notación Excel, 1-indexado):

- Columnas C/D, E/F, G/H, I/J, K/L son pares (PASADO, ACTUAL) para los
  5 tipos de cultura, en ese orden fijo: LOGRO, CENTRADA EN EL CLIENTE,
  EQUIPO UNICO, INNOVADORA, LAS PERSONAS PRIMERO.

- La hoja es una secuencia de "bloques". Cada bloque es:
      fila de ÍTEM      -> el texto de la frase evaluada, en C/E/G/I/K
      filas de RESPUESTA-> una por calificador (columna B = 1..12) y,
                           al final, la fila CONSENSO
      filas derivadas   -> eco del consenso (VLOOKUP), proporciones
                           R/A/V y VALORACIÓN; todas calculadas por
                           fórmulas, NO son respuestas

- La categoría (Tipo de cultura / Comportamientos / Símbolos / Sistemas)
  aparece como etiqueta en la columna B, pero **solo en el primer ítem de
  cada categoría**; los ítems siguientes la tienen vacía y la heredan.

Por eso los bloques NO se detectan por la etiqueta de la columna B (eso
perdía todos los ítems sin etiqueta), sino por las **rachas de filas de
respuesta**, que sí están siempre presentes en la plantilla: la fila de
ítem es la que está justo encima de cada racha.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from taller_cultura.application.ports import LectorTaller
from taller_cultura.application.validacion import Hallazgo, ResultadoValidacion, Severidad
from taller_cultura.domain.model import (
    Aspecto,
    Calificacion,
    Calificador,
    CategoriaAspecto,
    Momento,
    TipoCultura,
    Valoracion,
)

HOJA_CALIFICADORES = "CALIFICADORES"
HOJA_TALLER = "TALLER"

CODIGO_CALIFICADOR_CONSENSO = 0
NOMBRE_CALIFICADOR_CONSENSO = "CONSENSO"

# Columnas (1-indexadas, estilo Excel) de la hoja TALLER: tipo_cultura -> (col_pasado, col_actual)
COLUMNAS_TIPO_CULTURA: dict[TipoCultura, tuple[int, int]] = {
    TipoCultura.LOGRO: (3, 4),
    TipoCultura.CENTRADA_EN_EL_CLIENTE: (5, 6),
    TipoCultura.EQUIPO_UNICO: (7, 8),
    TipoCultura.INNOVADORA: (9, 10),
    TipoCultura.LAS_PERSONAS_PRIMERO: (11, 12),
}
COLUMNA_ETIQUETA_ITEM = 2  # columna B

# Etiquetas normalizadas (minúsculas, sin acentos) que fijan la categoría vigente.
ETIQUETA_A_CATEGORIA: dict[str, CategoriaAspecto] = {
    "tipo de cultura": CategoriaAspecto.TIPO_DE_CULTURA,
    "otras palabras": CategoriaAspecto.TIPO_DE_CULTURA,
    "una definicion": CategoriaAspecto.TIPO_DE_CULTURA,
    "comportamientos": CategoriaAspecto.COMPORTAMIENTOS,
    "simbolos": CategoriaAspecto.SIMBOLOS,
    "sistemas": CategoriaAspecto.SISTEMAS,
}

# Una racha de respuestas legítima tiene 7..13 filas. Este mínimo descarta
# los números sueltos de las filas de encabezado (fila 2 "CONSENSO", fila 4
# con los índices de columna 2/3/4/...), que si no se leerían como si
# fueran el inicio de un bloque.
MIN_FILAS_POR_RACHA = 3

# Cuántas filas sin identificar se toleran seguidas dentro de una racha antes
# de darla por terminada. Entre el último calificador de un bloque y el primero
# del siguiente hay al menos 6 filas (eco, 3 proporciones, VALORACIÓN, ítem),
# así que con 2 se salvan los huecos internos sin llegar a fundir dos bloques.
MAX_HUECOS_EN_RACHA = 2


def _normalizar(texto: str) -> str:
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return sin_acentos.strip().lower()


def _es_valoracion_valida(texto: object) -> bool:
    return isinstance(texto, str) and texto.strip().upper() in {"A", "R", "V"}


def _es_hueco_tolerable(valor: object) -> bool:
    """Fila sin identificar intercalada entre calificadores.

    Ocurre de dos formas: una fórmula rota (`#ERROR!`, `#N/A`) en el libro
    original, o una celda ya vacía si esa fórmula se limpió al preparar la
    plantilla. En ambos casos hay que saltarla sin dar por terminada la
    racha; si no, se pierde la fila CONSENSO que viene justo después.
    """
    if valor is None:
        return True
    return isinstance(valor, str) and (not valor.strip() or valor.strip().startswith("#"))


@dataclass
class _BloqueItem:
    """Un ítem calificable ya localizado en la hoja: la fila del texto, el
    rango de filas de respuesta y, para cada tipo de cultura presente, el
    Aspecto de dominio junto con el id temporal (posición en la lista
    plana) que lo identifica antes de que el repositorio le asigne un id.
    """

    fila_item: int
    fila_inicio_respuestas: int
    fila_fin_respuestas: int
    aspectos_por_cultura: dict[TipoCultura, tuple[int, Aspecto]]


class LectorTallerExcel(LectorTaller):
    """Lee empresas, calificadores, aspectos y calificaciones del `.xlsx`."""

    def __init__(self, ruta_excel: str | Path) -> None:
        self._ruta = Path(ruta_excel)
        self._workbook = openpyxl.load_workbook(self._ruta, data_only=True)
        self._bloques: list[_BloqueItem] | None = None

    def leer_nombre_empresa(self) -> str | None:
        """El libro rara vez trae el nombre; sirve solo para pre-rellenarlo."""
        if HOJA_CALIFICADORES not in self._workbook.sheetnames:
            return None
        hoja = self._workbook[HOJA_CALIFICADORES]
        nombre = hoja["C1"].value or hoja["B1"].value
        if not nombre or str(nombre).strip().upper() in {"EMPRESA", ""}:
            return None
        return str(nombre).strip()

    # -- validación previa --------------------------------------------------

    def validar(self) -> ResultadoValidacion:
        """Revisa que el libro sea realmente un taller procesable.

        Se comprueba en orden de gravedad: primero que existan las hojas y
        la estructura mínima (sin eso no hay nada que hacer), y después la
        riqueza de los datos (que se avisa, pero no impide procesar).
        """
        hallazgos: list[Hallazgo] = []

        if HOJA_TALLER not in self._workbook.sheetnames:
            hallazgos.append(
                Hallazgo(
                    Severidad.ERROR,
                    f"El libro no tiene la hoja «{HOJA_TALLER}». "
                    "¿Seguro que es el archivo del taller de cultura?",
                )
            )
            return ResultadoValidacion(hallazgos=tuple(hallazgos))

        if HOJA_CALIFICADORES not in self._workbook.sheetnames:
            hallazgos.append(
                Hallazgo(
                    Severidad.AVISO,
                    f"Falta la hoja «{HOJA_CALIFICADORES}»: no se podrán "
                    "nombrar los participantes, pero el taller sí se procesa.",
                )
            )

        hallazgos.extend(self._validar_encabezado_de_culturas())

        bloques = self._obtener_bloques()
        if not bloques:
            hallazgos.append(
                Hallazgo(
                    Severidad.ERROR,
                    "No se reconoció ningún ítem calificable en la hoja TALLER. "
                    "La estructura del archivo no coincide con la del taller.",
                )
            )
            return ResultadoValidacion(hallazgos=tuple(hallazgos))

        hallazgos.extend(self._validar_datos(bloques))
        return ResultadoValidacion(hallazgos=tuple(hallazgos))

    def _validar_encabezado_de_culturas(self) -> list[Hallazgo]:
        """La fila de encabezado debe nombrar los 5 tipos de cultura en sus
        columnas; si no, las columnas están corridas y todo saldría mal.
        """
        hoja = self._workbook[HOJA_TALLER]
        esperadas = {_normalizar(tc.value) for tc in TipoCultura}
        encontradas: set[str] = set()
        for fila in range(1, min(hoja.max_row, 30) + 1):
            for _tipo_cultura, (col_pasado, _) in COLUMNAS_TIPO_CULTURA.items():
                valor = hoja.cell(fila, col_pasado).value
                if isinstance(valor, str) and _normalizar(valor) in esperadas:
                    encontradas.add(_normalizar(valor))

        faltantes = esperadas - encontradas
        if not faltantes:
            return []
        if len(faltantes) == len(esperadas):
            return [
                Hallazgo(
                    Severidad.ERROR,
                    "En la hoja TALLER no aparece ninguno de los 5 tipos de cultura "
                    "en las columnas esperadas (C, E, G, I, K).",
                )
            ]
        return [
            Hallazgo(
                Severidad.AVISO,
                f"No se encontró el encabezado de {len(faltantes)} tipo(s) de cultura; "
                "se procesará igual, pero conviene revisar las columnas.",
            )
        ]

    def _validar_datos(self, bloques: list[_BloqueItem]) -> list[Hallazgo]:
        hallazgos: list[Hallazgo] = []
        aspectos = self.leer_aspectos()
        calificaciones = self.leer_calificaciones(aspectos)

        if not calificaciones:
            hallazgos.append(
                Hallazgo(
                    Severidad.AVISO,
                    "El taller está en blanco: no hay ninguna valoración registrada. "
                    "El reporte saldrá vacío.",
                )
            )
            return hallazgos

        consenso = [c for c in calificaciones if c.es_consenso]
        if not consenso:
            hallazgos.append(
                Hallazgo(
                    Severidad.AVISO,
                    "No hay filas de CONSENSO llenas. El reporte usará las "
                    "calificaciones individuales, que no es la base habitual del taller.",
                )
            )

        for momento in (Momento.PASADO, Momento.ACTUAL):
            if not any(c.momento is momento for c in (consenso or calificaciones)):
                hallazgos.append(
                    Hallazgo(
                        Severidad.AVISO,
                        f"El momento {momento.value} no tiene valoraciones: "
                        "no se podrá calcular la brecha entre ambos momentos.",
                    )
                )

        calificados = {c.aspecto_id for c in calificaciones}
        cobertura = 100 * len(calificados) / len(aspectos) if aspectos else 0
        if cobertura < 50:
            hallazgos.append(
                Hallazgo(
                    Severidad.AVISO,
                    f"Solo {cobertura:.0f}% de los ítems tiene alguna valoración; "
                    "los promedios se apoyan en pocos datos.",
                )
            )

        hallazgos.extend(self._validar_items_repetidos(aspectos))
        return hallazgos

    @staticmethod
    def _validar_items_repetidos(aspectos: list[Aspecto]) -> list[Hallazgo]:
        """Frases repetidas dentro de la misma cultura y categoría.

        Suelen venir de copiar y pegar filas al armar el taller. No impiden
        procesar, pero cuentan doble en los promedios, así que conviene
        saberlo antes de leer el reporte.
        """
        vistos: dict[tuple, int] = {}
        for aspecto in aspectos:
            clave = (aspecto.tipo_cultura, aspecto.categoria, _normalizar(aspecto.texto))
            vistos[clave] = vistos.get(clave, 0) + 1

        repetidos = {clave: n for clave, n in vistos.items() if n > 1}
        if not repetidos:
            return []

        # Se cuentan por frase, no por (frase × cultura), que es como lo ve
        # quien armó el taller.
        frases = {clave[2] for clave in repetidos}
        return [
            Hallazgo(
                Severidad.AVISO,
                f"Hay {len(frases)} ítem(s) repetido(s) en el taller (la misma frase "
                "aparece más de una vez en la misma categoría); esas valoraciones "
                "pesan doble en los promedios.",
            )
        ]

    def leer_calificadores(self) -> list[Calificador]:
        """Lee el roster de la hoja CALIFICADORES (filas 4 a 15: COD/NOMBRES)."""
        df = pd.read_excel(
            self._ruta,
            sheet_name=HOJA_CALIFICADORES,
            header=None,
            skiprows=3,
            nrows=12,
            usecols=[1, 2],
            names=["codigo", "nombre"],
        )
        calificadores = [
            Calificador(codigo=CODIGO_CALIFICADOR_CONSENSO, nombre=NOMBRE_CALIFICADOR_CONSENSO)
        ]
        for fila in df.itertuples(index=False):
            if pd.isna(fila.codigo):
                continue
            codigo = int(fila.codigo)
            nombre = str(fila.nombre) if not pd.isna(fila.nombre) else str(codigo)
            calificadores.append(Calificador(codigo=codigo, nombre=nombre))
        return calificadores

    def leer_aspectos(self) -> list[Aspecto]:
        aspectos: list[Aspecto] = []
        for bloque in self._obtener_bloques():
            for _tipo_cultura, (_id_temporal, aspecto) in bloque.aspectos_por_cultura.items():
                aspectos.append(aspecto)
        return aspectos

    def leer_calificaciones(self, aspectos: list[Aspecto]) -> list[Calificacion]:
        """El id temporal de cada calificación coincide con la posición del
        aspecto correspondiente en la lista que devuelve `leer_aspectos`,
        que la capa de aplicación remapea luego al id real del repositorio.
        """
        calificaciones: list[Calificacion] = []
        for bloque in self._obtener_bloques():
            calificaciones.extend(self._leer_respuestas_del_bloque(bloque))
        return calificaciones

    # -- localización de bloques -------------------------------------------

    def _obtener_bloques(self) -> list[_BloqueItem]:
        if self._bloques is None:
            self._bloques = list(self._localizar_bloques())
        return self._bloques

    def _localizar_bloques(self):
        hoja = self._workbook[HOJA_TALLER]
        max_row = hoja.max_row

        categoria_por_fila = self._mapear_categoria_vigente(hoja, max_row)
        orden_por_categoria: dict[CategoriaAspecto, int] = {}
        id_temporal = 0

        for inicio, fin in self._localizar_rachas_de_respuesta(hoja, max_row):
            fila_item = inicio - 1
            if fila_item < 1:
                continue
            categoria = categoria_por_fila.get(fila_item)
            if categoria is None:
                continue

            textos: dict[TipoCultura, str] = {}
            for tipo_cultura, (col_pasado, _col_actual) in COLUMNAS_TIPO_CULTURA.items():
                texto = hoja.cell(fila_item, col_pasado).value
                if isinstance(texto, str) and texto.strip():
                    textos[tipo_cultura] = texto.strip()
            if not textos:
                continue

            orden = orden_por_categoria.get(categoria, 0) + 1
            orden_por_categoria[categoria] = orden

            aspectos_por_cultura: dict[TipoCultura, tuple[int, Aspecto]] = {}
            for tipo_cultura, texto in textos.items():
                aspectos_por_cultura[tipo_cultura] = (
                    id_temporal,
                    Aspecto(
                        id=None,
                        tipo_cultura=tipo_cultura,
                        categoria=categoria,
                        orden=orden,
                        texto=texto,
                    ),
                )
                id_temporal += 1

            yield _BloqueItem(
                fila_item=fila_item,
                fila_inicio_respuestas=inicio,
                fila_fin_respuestas=fin,
                aspectos_por_cultura=aspectos_por_cultura,
            )

    @staticmethod
    def _mapear_categoria_vigente(hoja, max_row: int) -> dict[int, CategoriaAspecto | None]:
        """Para cada fila, qué categoría está vigente en ese punto de la hoja.

        La etiqueta solo aparece en el primer ítem de cada categoría; las
        filas siguientes la heredan hasta que aparezca la siguiente.
        """
        categoria_por_fila: dict[int, CategoriaAspecto | None] = {}
        categoria_actual: CategoriaAspecto | None = None
        for fila in range(1, max_row + 1):
            etiqueta = hoja.cell(fila, COLUMNA_ETIQUETA_ITEM).value
            if isinstance(etiqueta, str):
                clave = _normalizar(etiqueta)
                if clave in ETIQUETA_A_CATEGORIA:
                    categoria_actual = ETIQUETA_A_CATEGORIA[clave]
            categoria_por_fila[fila] = categoria_actual
        return categoria_por_fila

    def _localizar_rachas_de_respuesta(self, hoja, max_row: int) -> list[tuple[int, int]]:
        """Devuelve los rangos [inicio, fin] de filas de respuesta.

        Una racha son filas consecutivas cuya columna B identifica a un
        calificador (código 1..12) o a la fila CONSENSO. Las fórmulas rotas
        (`#ERROR!`) intercaladas se saltan sin cortar la racha. La racha
        termina en la primera fila que no es ninguna de las dos cosas —
        típicamente la fila de eco del consenso, que es un VLOOKUP y por lo
        tanto NO debe contarse como una respuesta más.
        """
        rachas: list[tuple[int, int]] = []
        fila = 1
        while fila <= max_row:
            if self._codigo_calificador_de_fila(hoja.cell(fila, COLUMNA_ETIQUETA_ITEM).value) is None:
                fila += 1
                continue

            inicio = fila
            fin = fila
            filas_de_respuesta = 0
            huecos_seguidos = 0
            cursor = fila
            while cursor <= max_row:
                valor_b = hoja.cell(cursor, COLUMNA_ETIQUETA_ITEM).value
                codigo = self._codigo_calificador_de_fila(valor_b)

                if codigo is not None:
                    fin = cursor
                    filas_de_respuesta += 1
                    huecos_seguidos = 0
                    cursor += 1
                    if codigo == CODIGO_CALIFICADOR_CONSENSO:
                        # El consenso cierra el bloque: lo que sigue (el eco
                        # del VLOOKUP, las proporciones y la VALORACIÓN) es
                        # cálculo derivado, no respuestas.
                        break
                    continue

                if _es_hueco_tolerable(valor_b) and huecos_seguidos < MAX_HUECOS_EN_RACHA:
                    huecos_seguidos += 1
                    cursor += 1
                    continue

                break

            if filas_de_respuesta >= MIN_FILAS_POR_RACHA:
                rachas.append((inicio, fin))
            fila = cursor if cursor > fila else fila + 1
        return rachas

    # -- lectura de respuestas ---------------------------------------------

    def _leer_respuestas_del_bloque(self, bloque: _BloqueItem) -> list[Calificacion]:
        hoja = self._workbook[HOJA_TALLER]
        calificaciones: list[Calificacion] = []

        for fila in range(bloque.fila_inicio_respuestas, bloque.fila_fin_respuestas + 1):
            codigo_calificador = self._codigo_calificador_de_fila(
                hoja.cell(fila, COLUMNA_ETIQUETA_ITEM).value
            )
            if codigo_calificador is None:
                continue  # fila con fórmula rota intercalada

            for tipo_cultura, (id_temporal, _aspecto) in bloque.aspectos_por_cultura.items():
                col_pasado, col_actual = COLUMNAS_TIPO_CULTURA[tipo_cultura]
                for momento, columna in (
                    (Momento.PASADO, col_pasado),
                    (Momento.ACTUAL, col_actual),
                ):
                    valor = hoja.cell(fila, columna).value
                    if not _es_valoracion_valida(valor):
                        continue
                    calificaciones.append(
                        Calificacion(
                            aspecto_id=id_temporal,
                            calificador_codigo=codigo_calificador,
                            momento=momento,
                            valoracion=Valoracion.desde_texto(valor),
                            es_consenso=(codigo_calificador == CODIGO_CALIFICADOR_CONSENSO),
                        )
                    )

        return calificaciones

    @staticmethod
    def _codigo_calificador_de_fila(etiqueta: object) -> int | None:
        if isinstance(etiqueta, str) and etiqueta.strip().upper() == NOMBRE_CALIFICADOR_CONSENSO:
            return CODIGO_CALIFICADOR_CONSENSO
        if isinstance(etiqueta, bool):
            return None
        if isinstance(etiqueta, (int, float)) and float(etiqueta).is_integer():
            return int(etiqueta)
        return None
