"""Adaptador de salida: produce el archivo que se reparte a los participantes.

Implementa el puerto `ExportadorPlantilla`. Del libro original conserva
únicamente las hojas de trabajo (`TALLER` y el roster `CALIFICADORES`) y
descarta todo el material derivado (`RESUMEN`, `PRESENTACION`,
`PRESENTACION 2`, `FINAL`, `CULTURAS`, `Hoja1`), cuyo contenido pasa a
vivir en el reporte que genera esta aplicación.

Las fórmulas se sustituyen por sus valores calculados: al quitar las hojas
derivadas, cualquier fórmula que las referenciara quedaría rota (`#REF!`),
y las de la propia hoja TALLER que apuntan a `CALIFICADORES` se conservan
correctas al escribirse ya resueltas.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HOJAS_QUE_SE_ENVIAN = ("CALIFICADORES", "TALLER")

# Restos de fórmulas rotas que no tiene sentido repartir; se dejan en blanco.
VALORES_DE_ERROR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#ERROR!"}


class ExportadorPlantillaExcel:
    """Escribe un `.xlsx` con solo las hojas de trabajo del taller."""

    def exportar_plantilla(self, ruta_excel_origen: str, ruta_destino: str) -> None:
        origen = openpyxl.load_workbook(ruta_excel_origen, data_only=True)
        destino = openpyxl.Workbook()
        destino.remove(destino.active)

        for nombre in HOJAS_QUE_SE_ENVIAN:
            if nombre not in origen.sheetnames:
                continue
            self._copiar_hoja(origen[nombre], destino.create_sheet(nombre))

        ruta = Path(ruta_destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        destino.save(ruta)

    @staticmethod
    def _copiar_hoja(hoja_origen, hoja_destino) -> None:
        for fila in hoja_origen.iter_rows():
            for celda in fila:
                valor = celda.value
                if isinstance(valor, str) and valor.strip() in VALORES_DE_ERROR:
                    valor = None
                if valor is None:
                    continue
                hoja_destino.cell(celda.row, celda.column, valor)

        for rango in list(hoja_origen.merged_cells.ranges):
            hoja_destino.merge_cells(str(rango))

        for letra, dim in hoja_origen.column_dimensions.items():
            if dim.width:
                hoja_destino.column_dimensions[letra].width = dim.width
        for indice, dim in hoja_origen.row_dimensions.items():
            if dim.height:
                hoja_destino.row_dimensions[indice].height = dim.height

        # Las columnas auxiliares de cálculo (M en adelante) no aportan nada
        # a quien llena el taller; se ocultan para dejar la hoja limpia.
        if hoja_destino.title == "TALLER":
            for indice_columna in range(13, 38):
                hoja_destino.column_dimensions[get_column_letter(indice_columna)].hidden = True
