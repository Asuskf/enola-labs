"""Test del exportador de la plantilla que se reparte.

La regla de negocio: del libro solo se envían las hojas de trabajo; todo
el material derivado (RESUMEN, PRESENTACION, FINAL, CULTURAS) se queda
fuera porque su contenido va en el reporte.
"""

from pathlib import Path

import openpyxl
import pytest

from taller_cultura.infrastructure.adapters.plantilla_writer import (
    HOJAS_QUE_SE_ENVIAN,
    VALORES_DE_ERROR,
    ExportadorPlantillaExcel,
)

RUTA_EXCEL = Path(__file__).resolve().parents[1] / "data" / "Taller para pasar a David (1).xlsx"

pytestmark = pytest.mark.skipif(
    not RUTA_EXCEL.exists(), reason="Excel del taller no disponible en esta máquina"
)


@pytest.fixture(scope="module")
def plantilla(tmp_path_factory) -> Path:
    destino = tmp_path_factory.mktemp("plantilla") / "TALLER.xlsx"
    ExportadorPlantillaExcel().exportar_plantilla(str(RUTA_EXCEL), str(destino))
    return destino


def test_solo_conserva_las_hojas_de_trabajo(plantilla):
    hojas = openpyxl.load_workbook(plantilla).sheetnames

    assert set(hojas) == set(HOJAS_QUE_SE_ENVIAN)
    for derivada in ("RESUMEN", "FINAL", "CULTURAS", "PRESENTACION"):
        assert derivada not in hojas


def test_no_deja_formulas_ni_errores(plantilla):
    hoja = openpyxl.load_workbook(plantilla)["TALLER"]

    for fila in hoja.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str):
                assert not celda.value.startswith("="), f"{celda.coordinate} quedó como fórmula"
                assert not celda.value.startswith("#"), f"{celda.coordinate} quedó con error"


def test_conserva_el_contenido_del_taller(plantilla):
    """No se pierde ninguna celda con contenido aprovechable.

    Dos matices deliberados:
    - se compara contra las celdas con valor, no contra `max_row`: el libro
      original arrastra cientos de filas vacías con formato que openpyxl
      cuenta como usadas y que no aportan nada;
    - los valores de error (`#N/A`, `#REF!`…) quedan fuera porque el
      exportador los limpia a propósito: son fórmulas rotas de la plantilla
      original y no tiene sentido repartirlas.
    """
    original = openpyxl.load_workbook(RUTA_EXCEL, data_only=True)["TALLER"]
    copia = openpyxl.load_workbook(plantilla)["TALLER"]

    def celdas_con_valor(hoja):
        return {
            (c.row, c.column)
            for fila in hoja.iter_rows()
            for c in fila
            if c.value is not None
            and str(c.value).strip() != ""
            and str(c.value).strip() not in VALORES_DE_ERROR
        }

    faltantes = celdas_con_valor(original) - celdas_con_valor(copia)
    assert not faltantes, f"se perdieron {len(faltantes)} celdas al exportar"

    # El roster resuelto y las filas clave siguen ahí.
    assert copia["B7"].value == 1
    assert str(copia["B45"].value).strip().upper() == "CONSENSO"
